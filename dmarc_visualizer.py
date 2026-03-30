#!/usr/bin/env python3
"""
DMARC Report Visualizer — v4
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
Uso:
    python3 dmarc_visualizer.py                        # carpeta: ./dmarc_reports
    python3 dmarc_visualizer.py /ruta/reportes
    python3 dmarc_visualizer.py /ruta out.html --no-open

Requisitos:
    pip install openpyxl
"""

import os, sys, xml.etree.ElementTree as ET, json, zipfile, gzip
import webbrowser, http.server, threading, socket, time, hashlib
from datetime import datetime, timezone
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False
    print("  ⚠  openpyxl no encontrado. Instala: pip install openpyxl")


# ══════════════════════════════════════════════════════════════════
#  CONSTANTS
# ══════════════════════════════════════════════════════════════════

EXCEL_FILE   = "dmarc_database.xlsx"
VERSION      = "4.0"


# ══════════════════════════════════════════════════════════════════
#  EXCEL — helpers
# ══════════════════════════════════════════════════════════════════

HDR_FILL  = PatternFill("solid", fgColor="0f172a") if EXCEL_OK else None
HDR_FONT  = Font(bold=True, color="22d3ee", size=10) if EXCEL_OK else None
EVEN_FILL = PatternFill("solid", fgColor="f8fafc") if EXCEL_OK else None

def _hdr(ws, row, ncols):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26

def _zebra(cell, row):
    if row % 2 == 0:
        cell.fill = EVEN_FILL

def _widths(ws, ws_list):
    for i, w in enumerate(ws_list, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def _write_row(ws, row_idx, values, color_map=None):
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col, value=val)
        _zebra(cell, row_idx)
        if color_map and col in color_map:
            cell.font = Font(color=color_map[col])
    return row_idx + 1


# ══════════════════════════════════════════════════════════════════
#  EXCEL — init / load
# ══════════════════════════════════════════════════════════════════

def init_excel(folder):
    """
    Load existing Excel DB or create fresh one.
    Returns (wb, ignored_ips:set, ip_labels:dict, seen_hashes:set)
    """
    path = Path(folder) / EXCEL_FILE
    ignored_ips  = set()
    ip_labels    = {}
    seen_hashes  = set()

    if not EXCEL_OK:
        return None, ignored_ips, ip_labels, seen_hashes

    if path.exists():
        try:
            wb = openpyxl.load_workbook(path)

            # IPs ignoradas
            if "IPs Ignoradas" in wb.sheetnames:
                for row in wb["IPs Ignoradas"].iter_rows(min_row=2, values_only=True):
                    if row and row[0]:
                        ip = str(row[0]).strip()
                        ignored_ips.add(ip)
                        if len(row) > 1 and row[1]:
                            ip_labels[ip] = str(row[1]).strip()

            # Hashes de reportes ya procesados
            if "Historial Reportes" in wb.sheetnames:
                for row in wb["Historial Reportes"].iter_rows(min_row=2, values_only=True):
                    if row and row[0]:
                        seen_hashes.add(str(row[0]).strip())

            n_ign = len(ignored_ips)
            n_hsh = len(seen_hashes)
            print(f"  📊 Excel cargado: {path.name}  ({n_ign} IPs ignoradas, {n_hsh} reportes en historial)")
            return wb, ignored_ips, ip_labels, seen_hashes
        except Exception as e:
            print(f"  ⚠  Error leyendo Excel: {e} — se creará uno nuevo")

    # ── Crear workbook nuevo ──────────────────────────────────────
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Hoja 1 – Resumen por reporte
    ws = wb.create_sheet("Resumen")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "DMARC DATABASE"
    ws["A1"].font = Font(bold=True, size=14, color="0f172a")
    ws["A2"] = f"Creado: {datetime.now().strftime('%Y-%m-%d %H:%M')}  —  v{VERSION}"
    ws["A2"].font = Font(size=9, color="64748b")
    headers = ["Archivo","Organización","Dominio","Fecha inicio","Fecha fin",
               "Total msgs","DMARC Pass","DMARC Fail","Pass Rate %",
               "Solo DKIM","Solo SPF","Cuarentena","Rechazados","IPs únicas","Registros"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    _hdr(ws, 4, len(headers))
    _widths(ws, [32,22,18,20,20,12,12,12,12,11,11,12,12,12,10])

    # Hoja 2 – Todos los registros
    ws2 = wb.create_sheet("Registros")
    ws2.sheet_view.showGridLines = False
    h2 = ["Hash Reporte","Archivo","Org","Dominio","Fecha","IP Origen","Mensajes",
          "Header From","DKIM eval","SPF eval","DMARC result","Disposición",
          "Alineación","Auth DKIM Domain","DKIM Selector","Auth SPF Domain",
          "Envelope From","Razones"]
    for i, h in enumerate(h2, 1): ws2.cell(row=1, column=i, value=h)
    _hdr(ws2, 1, len(h2))
    _widths(ws2, [14,30,18,16,12,22,10,22,10,10,12,12,14,24,28,22,22,35])

    # Hoja 3 – IPs Ignoradas (EDITABLE por el usuario)
    ws3 = wb.create_sheet("IPs Ignoradas")
    ws3.sheet_view.showGridLines = False
    ws3["A1"] = "IP / Rango"
    ws3["B1"] = "Etiqueta / Descripción"
    ws3["C1"] = "Fecha añadida"
    ws3["D1"] = "Motivo"
    _hdr(ws3, 1, 4)
    _widths(ws3, [24, 38, 18, 45])
    # Instrucciones
    ws3["A2"] = "← Escribe aquí las IPs a ignorar"
    ws3["A2"].font = Font(italic=True, color="94a3b8", size=9)
    ws3.merge_cells("A2:D2")

    # Hoja 4 – Estadísticas por IP
    ws4 = wb.create_sheet("Estadísticas IPs")
    ws4.sheet_view.showGridLines = False
    h4 = ["IP","Total msgs","DMARC Pass","DMARC Fail","Cuarentena",
          "Rechazados","% Pass","Dominios vistos","Etiqueta","Ignorada"]
    for i, h in enumerate(h4, 1): ws4.cell(row=1, column=i, value=h)
    _hdr(ws4, 1, len(h4))
    _widths(ws4, [22,13,13,13,13,12,10,35,30,10])

    # Hoja 5 – Historial de reportes procesados
    ws5 = wb.create_sheet("Historial Reportes")
    ws5.sheet_view.showGridLines = False
    ws5["A1"] = "Hash (SHA256)"
    ws5["B1"] = "Archivo"
    ws5["C1"] = "Organización"
    ws5["D1"] = "Dominio"
    ws5["E1"] = "Fecha reporte"
    ws5["F1"] = "Fecha procesado"
    ws5["G1"] = "Total msgs"
    ws5["H1"] = "DMARC Pass %"
    _hdr(ws5, 1, 8)
    _widths(ws5, [20,32,22,18,20,20,12,13])

    # Hoja 6 – Dominios
    ws6 = wb.create_sheet("Dominios")
    ws6.sheet_view.showGridLines = False
    h6 = ["Dominio","Total msgs","DMARC Pass","DMARC Fail","Pass Rate %",
          "Cuarentena","Rechazados","Orgs reportando","Último reporte"]
    for i, h in enumerate(h6, 1): ws6.cell(row=1, column=i, value=h)
    _hdr(ws6, 1, len(h6))
    _widths(ws6, [22,13,13,13,12,13,13,25,20])

    wb.save(path)
    print(f"  📊 Excel nuevo creado: {path.name}")
    return wb, ignored_ips, ip_labels, seen_hashes


def update_excel(wb, reports, stats, folder, new_reports_only):
    """Write data into the Excel workbook and save."""
    if not EXCEL_OK or wb is None:
        return
    path = Path(folder) / EXCEL_FILE

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")

    # ── Hoja 1: Resumen ─────────────────────────────────────────
    ws1 = wb["Resumen"]
    ws1["A2"] = f"Actualizado: {now_str}  —  {len(reports)} reportes  —  v{VERSION}"

    # Clear old data rows
    max_row = ws1.max_row
    for r in range(5, max_row + 1):
        for c in range(1, 16):
            ws1.cell(row=r, column=c).value = None

    ri = 5
    for rpt in reports:
        s = rpt.get("summary", {})
        pr = s.get("dmarc_pass_rate", 0)
        vals = [
            rpt.get("filename",""), rpt.get("org_name",""), rpt.get("domain",""),
            rpt.get("date_begin",""), rpt.get("date_end",""),
            s.get("total_messages",0), s.get("dmarc_pass",0), s.get("dmarc_fail",0),
            pr/100,
            s.get("dkim_only",0), s.get("spf_only",0),
            s.get("quarantined",0), s.get("rejected",0),
            s.get("unique_ips",0), s.get("record_count",0),
        ]
        for ci, val in enumerate(vals, 1):
            cell = ws1.cell(row=ri, column=ci, value=val)
            _zebra(cell, ri)
            if ci == 9:
                cell.number_format = "0.0%"
                color = "15803d" if pr >= 90 else "854d0e" if pr >= 70 else "991b1b"
                cell.font = Font(color=color, bold=True)
        ri += 1

    # ── Hoja 2: Registros ────────────────────────────────────────
    ws2 = wb["Registros"]
    # Append only new reports
    next_row = ws2.max_row + 1
    if next_row == 2 and ws2.cell(2, 1).value is None:
        next_row = 2

    existing_hashes = set()
    for row in ws2.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]: existing_hashes.add(str(row[0]))

    for rpt in reports:
        rpt_hash = rpt.get("_hash", "")
        if rpt_hash in existing_hashes:
            continue  # already written
        for rec in rpt.get("records", []):
            reasons = "; ".join(f"{x.get('type','')}:{x.get('comment','')}" for x in rec.get("reasons",[]) if x.get("type"))
            vals = [
                rpt_hash, rpt.get("filename",""), rpt.get("org_name",""),
                rpt.get("domain",""), rpt.get("date_label",""),
                rec.get("source_ip",""), rec.get("count",0),
                rec.get("header_from",""),
                rec.get("dkim_result",""), rec.get("spf_result",""),
                rec.get("dmarc_result",""), rec.get("disposition",""),
                rec.get("alignment",""), rec.get("auth_dkim_domain",""),
                rec.get("auth_dkim_selector",""), rec.get("auth_spf_domain",""),
                rec.get("envelope_from",""), reasons,
            ]
            dkim_col  = {9: "15803d" if rec.get("dkim_result")=="pass" else "991b1b"}
            spf_col   = {10: "15803d" if rec.get("spf_result")=="pass" else "991b1b"}
            dmarc_col = {11: "15803d" if rec.get("dmarc_result")=="pass" else "991b1b"}
            cmap = {**dkim_col, **spf_col, **dmarc_col}
            for ci, val in enumerate(vals, 1):
                cell = ws2.cell(row=next_row, column=ci, value=val)
                _zebra(cell, next_row)
                if ci in cmap:
                    cell.font = Font(color=cmap[ci])
            next_row += 1

    # ── Hoja 4: Estadísticas IPs ─────────────────────────────────
    ws4 = wb["Estadísticas IPs"]
    for r in range(2, ws4.max_row + 1):
        for c in range(1, 11): ws4.cell(row=r, column=c).value = None

    ip_domains = defaultdict(set)
    for rpt in reports:
        dom = rpt.get("domain","")
        for rec in rpt.get("records",[]):
            ip_domains[rec.get("source_ip","")].add(dom)

    ri = 2
    for ip, total in sorted(stats.get("sources",{}).items(), key=lambda x: -x[1]):
        p   = stats.get("ip_pass",{}).get(ip,0)
        f   = stats.get("ip_fail",{}).get(ip,0)
        q   = stats.get("ip_quarantine",{}).get(ip,0)
        rej = stats.get("ip_reject",{}).get(ip,0)
        pr  = round(p/total*100,1) if total else 0
        doms = ", ".join(sorted(ip_domains[ip]))
        lbl  = stats.get("ip_labels",{}).get(ip,"")
        ign  = "Sí" if ip in stats.get("ignored_ips",[]) else "No"
        vals = [ip, total, p, f, q, rej, pr/100, doms, lbl, ign]
        for ci, val in enumerate(vals, 1):
            cell = ws4.cell(row=ri, column=ci, value=val)
            _zebra(cell, ri)
            if ci == 7:
                cell.number_format = "0.0%"
                color = "15803d" if pr >= 90 else "854d0e" if pr >= 70 else "991b1b"
                cell.font = Font(color=color, bold=True)
        ri += 1

    # ── Hoja 5: Historial ────────────────────────────────────────
    ws5 = wb["Historial Reportes"]
    existing_h = set()
    for row in ws5.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]: existing_h.add(str(row[0]))

    hr = ws5.max_row + 1
    if hr == 2 and ws5.cell(2,1).value is None: hr = 2
    for rpt in reports:
        h = rpt.get("_hash","")
        if h in existing_h: continue
        s = rpt.get("summary",{})
        vals = [h, rpt.get("filename",""), rpt.get("org_name",""),
                rpt.get("domain",""), rpt.get("date_begin",""),
                now_str, s.get("total_messages",0),
                s.get("dmarc_pass_rate",0)/100]
        for ci, val in enumerate(vals, 1):
            cell = ws5.cell(row=hr, column=ci, value=val)
            _zebra(cell, hr)
            if ci == 8: cell.number_format = "0.0%"
        hr += 1

    # ── Hoja 6: Dominios ─────────────────────────────────────────
    ws6 = wb["Dominios"]
    for r in range(2, ws6.max_row + 1):
        for c in range(1, 10): ws6.cell(row=r, column=c).value = None

    dom_stats = defaultdict(lambda: {"msgs":0,"pass":0,"fail":0,"quar":0,"rej":0,"orgs":set(),"last":""})
    for rpt in reports:
        dom = rpt.get("domain","")
        s   = rpt.get("summary",{})
        ds  = dom_stats[dom]
        ds["msgs"] += s.get("total_messages",0)
        ds["pass"] += s.get("dmarc_pass",0)
        ds["fail"] += s.get("dmarc_fail",0)
        ds["quar"] += s.get("quarantined",0)
        ds["rej"]  += s.get("rejected",0)
        ds["orgs"].add(rpt.get("org_name",""))
        if rpt.get("date_begin","") > ds["last"]:
            ds["last"] = rpt.get("date_begin","")

    ri = 2
    for dom, ds in sorted(dom_stats.items()):
        pr = round(ds["pass"]/ds["msgs"]*100,1) if ds["msgs"] else 0
        vals = [dom, ds["msgs"], ds["pass"], ds["fail"], pr/100,
                ds["quar"], ds["rej"], ", ".join(sorted(ds["orgs"])), ds["last"]]
        for ci, val in enumerate(vals, 1):
            cell = ws6.cell(row=ri, column=ci, value=val)
            _zebra(cell, ri)
            if ci == 5: cell.number_format = "0.0%"
        ri += 1

    wb.save(path)
    print(f"  💾 Excel actualizado: {path.name}")


# ══════════════════════════════════════════════════════════════════
#  PARSING
# ══════════════════════════════════════════════════════════════════

def _ts(ts_int):
    """Convert unix timestamp to UTC datetime string, handling future timestamps gracefully."""
    try:
        ts = int(ts_int)
        # Some providers send milliseconds instead of seconds
        if ts > 9_999_999_999:
            ts = ts // 1000
        dt = datetime.fromtimestamp(ts, tz=timezone.utc)
        return dt.strftime("%Y-%m-%d %H:%M UTC"), dt.strftime("%Y-%m-%d")
    except Exception:
        return "Unknown", "Unknown"


def parse_dmarc_xml(xml_content, filename=""):
    try:
        root = ET.fromstring(xml_content)
    except ET.ParseError as e:
        print(f"    ✗ XML error en {filename}: {e}")
        return None

    # Hash for deduplication
    file_hash = hashlib.sha256(xml_content.encode("utf-8", errors="replace")).hexdigest()[:16]

    report = {"filename": filename, "_hash": file_hash}

    # ── Metadata ─────────────────────────────────────────────────
    meta = root.find("report_metadata")
    if meta is not None:
        report["org_name"]      = (meta.findtext("org_name") or "Unknown").strip()
        report["email"]         = (meta.findtext("email") or "").strip()
        report["extra_contact"] = (meta.findtext("extra_contact_info") or "").strip()
        report["report_id"]     = (meta.findtext("report_id") or "N/A").strip()
        dr = meta.find("date_range")
        if dr is not None:
            b_str = dr.findtext("begin", "0")
            e_str = dr.findtext("end",   "0")
            report["date_begin"], report["date_label"] = _ts(b_str)
            report["date_end"],   _                    = _ts(e_str)
            try:
                report["date_begin_ts"] = int(b_str)
            except Exception:
                report["date_begin_ts"] = 0
    report.setdefault("org_name", "Unknown")
    report.setdefault("report_id", "N/A")
    report.setdefault("date_begin", "Unknown")
    report.setdefault("date_end",   "Unknown")
    report.setdefault("date_label", "Unknown")

    # ── Policy ───────────────────────────────────────────────────
    pol = root.find("policy_published")
    if pol is not None:
        report["domain"]       = (pol.findtext("domain") or "Unknown").strip()
        report["policy_p"]     = (pol.findtext("p")    or "none").strip()
        report["policy_sp"]    = (pol.findtext("sp")   or "none").strip()
        report["policy_np"]    = (pol.findtext("np")   or "").strip()
        report["policy_adkim"] = (pol.findtext("adkim") or "r").strip()
        report["policy_aspf"]  = (pol.findtext("aspf")  or "r").strip()
        report["policy_pct"]   = (pol.findtext("pct")   or "100").strip()
    report.setdefault("domain", "Unknown")

    # ── Records ──────────────────────────────────────────────────
    records = []
    for record in root.findall("record"):
        rec = {}
        row = record.find("row")
        if row is not None:
            rec["source_ip"] = (row.findtext("source_ip") or "").strip()
            rec["count"]     = int(row.findtext("count", "0") or 0)
            pe = row.find("policy_evaluated")
            if pe is not None:
                rec["disposition"] = (pe.findtext("disposition") or "none").strip()
                rec["dkim_result"] = (pe.findtext("dkim") or "fail").strip()
                rec["spf_result"]  = (pe.findtext("spf")  or "fail").strip()
                rec["reasons"]     = [
                    {"type":    (r.findtext("type")    or "").strip(),
                     "comment": (r.findtext("comment") or "").strip()}
                    for r in pe.findall("reason")
                ]
            else:
                rec["disposition"] = "none"; rec["dkim_result"] = "unknown"
                rec["spf_result"]  = "unknown"; rec["reasons"] = []
        else:
            rec = {"source_ip":"","count":0,"disposition":"none",
                   "dkim_result":"unknown","spf_result":"unknown","reasons":[]}

        ident = record.find("identifiers")
        if ident is not None:
            rec["header_from"]   = (ident.findtext("header_from")   or "").strip()
            rec["envelope_from"] = (ident.findtext("envelope_from") or "").strip()
        else:
            rec["header_from"] = rec["envelope_from"] = ""

        auth = record.find("auth_results")
        rec["auth_dkim_entries"] = []
        rec["auth_spf_entries"]  = []
        if auth is not None:
            for el in auth.findall("dkim"):
                rec["auth_dkim_entries"].append({
                    "domain":   (el.findtext("domain")   or "").strip(),
                    "result":   (el.findtext("result")   or "").strip(),
                    "selector": (el.findtext("selector") or "").strip(),
                })
            for el in auth.findall("spf"):
                rec["auth_spf_entries"].append({
                    "domain": (el.findtext("domain") or "").strip(),
                    "result": (el.findtext("result") or "").strip(),
                })

        rec["auth_dkim_domain"]   = rec["auth_dkim_entries"][0]["domain"]   if rec["auth_dkim_entries"] else ""
        rec["auth_dkim_result"]   = rec["auth_dkim_entries"][0]["result"]   if rec["auth_dkim_entries"] else ""
        rec["auth_dkim_selector"] = rec["auth_dkim_entries"][0]["selector"] if rec["auth_dkim_entries"] else ""
        rec["auth_spf_domain"]    = rec["auth_spf_entries"][0]["domain"]    if rec["auth_spf_entries"] else ""
        rec["auth_spf_result"]    = rec["auth_spf_entries"][0]["result"]    if rec["auth_spf_entries"] else ""

        d, s = rec["dkim_result"], rec["spf_result"]
        # DMARC pass = EITHER dkim OR spf aligns (per RFC 7489)
        rec["dmarc_result"] = "pass" if (d == "pass" or s == "pass") else "fail"
        rec["alignment"] = ("full_pass" if d == "pass" and s == "pass"
                            else "dkim_only" if d == "pass"
                            else "spf_only"  if s == "pass"
                            else "fail")
        records.append(rec)

    report["records"] = records

    total      = sum(r["count"] for r in records)
    dmarc_pass = sum(r["count"] for r in records if r["dmarc_result"] == "pass")
    dmarc_fail = total - dmarc_pass

    report["summary"] = {
        "total_messages":  total,
        "dmarc_pass":      dmarc_pass,
        "dmarc_fail":      dmarc_fail,
        "full_pass":       sum(r["count"] for r in records if r["alignment"] == "full_pass"),
        "dkim_only":       sum(r["count"] for r in records if r["alignment"] == "dkim_only"),
        "spf_only":        sum(r["count"] for r in records if r["alignment"] == "spf_only"),
        "align_fail":      sum(r["count"] for r in records if r["alignment"] == "fail"),
        "quarantined":     sum(r["count"] for r in records if r["disposition"] == "quarantine"),
        "rejected":        sum(r["count"] for r in records if r["disposition"] == "reject"),
        "dmarc_pass_rate": round(dmarc_pass / total * 100, 1) if total > 0 else 0,
        "unique_ips":      len(set(r["source_ip"] for r in records)),
        "record_count":    len(records),
    }
    return report


# ══════════════════════════════════════════════════════════════════
#  FILE LOADING
# ══════════════════════════════════════════════════════════════════

def read_file(path):
    path = Path(path)
    ext  = path.suffix.lower()
    try:
        if ext == ".zip":
            with zipfile.ZipFile(path, "r") as zf:
                for name in zf.namelist():
                    if name.lower().endswith(".xml"):
                        with zf.open(name) as f:
                            return f.read().decode("utf-8", errors="replace")
                if zf.namelist():
                    with zf.open(zf.namelist()[0]) as f:
                        return f.read().decode("utf-8", errors="replace")
        elif ext == ".gz":
            with gzip.open(path, "rb") as f:
                return f.read().decode("utf-8", errors="replace")
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
    except Exception as e:
        print(f"    ✗ No se puede leer {path.name}: {e}")
    return None


def load_folder(folder, seen_hashes=None):
    folder      = Path(folder)
    seen_hashes = seen_hashes or set()

    if not folder.exists():
        print(f"\n  ✗ Carpeta no encontrada: {folder}")
        folder.mkdir(parents=True, exist_ok=True)
        return [], []

    files = sorted({
        f for p in ["*.xml","*.zip","*.gz"]
        for f in folder.glob(p)
        if f.name != EXCEL_FILE
    })

    print(f"\n  📂 Carpeta   : {folder.resolve()}")
    print(f"  📄 Archivos  : {len(files)}\n")
    if not files:
        print("  ⚠  Sin archivos .xml / .zip / .gz")
        return [], []

    all_reports  = []
    new_reports  = []

    for path in files:
        size_kb = path.stat().st_size / 1024
        print(f"  ▸ {path.name:<52} ({size_kb:6.1f} KB) ... ", end="", flush=True)
        content = read_file(path)
        if content is None:
            print("SKIP"); continue

        report = parse_dmarc_xml(content, filename=path.name)
        if not report:
            print("ERROR"); continue

        s    = report["summary"]
        is_new = report["_hash"] not in seen_hashes
        tag  = "[NEW] " if is_new else "[ya procesado] "
        ok   = s["quarantined"]==0 and s["rejected"]==0 and s["dmarc_fail"]==0
        icon = "✓" if ok else "⚠"
        print(f"{icon} {tag}{s['record_count']} recs · {s['total_messages']} msgs · {s['dmarc_pass_rate']}% DMARC pass")

        all_reports.append(report)
        if is_new:
            new_reports.append(report)

    print(f"\n  ✅ {len(all_reports)} reportes  ({len(new_reports)} nuevos)\n")
    return all_reports, new_reports


# ══════════════════════════════════════════════════════════════════
#  GLOBAL STATS
# ══════════════════════════════════════════════════════════════════

def compute_stats(reports, ignored_ips=None, ip_labels=None):
    ignored_ips = ignored_ips or set()
    ip_labels   = ip_labels   or {}

    s = {
        "total_reports":  len(reports),
        "total_messages": 0,
        "dmarc_pass": 0, "dmarc_fail": 0,
        "full_pass": 0, "dkim_only": 0, "spf_only": 0, "fail": 0,
        "quarantined": 0, "rejected": 0,
        "sources":       defaultdict(int),
        "ip_pass":       defaultdict(int),
        "ip_fail":       defaultdict(int),
        "ip_quarantine": defaultdict(int),
        "ip_reject":     defaultdict(int),
        "ip_labels":     ip_labels,
        "domains":       defaultdict(lambda: {"msgs":0,"pass":0,"fail":0,"quar":0,"rej":0}),
        "orgs":          defaultdict(int),
        "timeline":      defaultdict(lambda: {"total":0,"pass":0,"fail":0,"quarantine":0}),
        "dispositions":  defaultdict(int),
        "alignment_breakdown": defaultdict(int),
        "quarantine_records": [],
        "reject_records":     [],
        "fail_records":       [],
        "spf_error_records":  [],
        "ignored_ips":   list(ignored_ips),
    }

    for rpt in reports:
        org    = rpt.get("org_name","?")
        domain = rpt.get("domain","?")
        date   = rpt.get("date_label","")
        s["orgs"][org] += 1
        s["domains"][domain]["msgs"] += rpt["summary"].get("total_messages",0)
        s["domains"][domain]["pass"] += rpt["summary"].get("dmarc_pass",0)
        s["domains"][domain]["fail"] += rpt["summary"].get("dmarc_fail",0)
        s["domains"][domain]["quar"] += rpt["summary"].get("quarantined",0)
        s["domains"][domain]["rej"]  += rpt["summary"].get("rejected",0)

        for rec in rpt.get("records",[]):
            n    = rec.get("count",0)
            aln  = rec.get("alignment","fail")
            dres = rec.get("dmarc_result","fail")
            disp = rec.get("disposition","none")
            ip   = rec.get("source_ip","?")

            s["total_messages"] += n
            s["dmarc_pass" if dres=="pass" else "dmarc_fail"] += n
            s[aln if aln in s else "fail"] += n
            s["dispositions"][disp] += n
            s["alignment_breakdown"][aln] += n
            s["sources"][ip] += n
            s["ip_pass"][ip]      += n if dres=="pass"      else 0
            s["ip_fail"][ip]      += n if dres=="fail"      else 0
            s["ip_quarantine"][ip]+= n if disp=="quarantine" else 0
            s["ip_reject"][ip]    += n if disp=="reject"     else 0

            if date:
                s["timeline"][date]["total"] += n
                s["timeline"][date]["pass" if dres=="pass" else "fail"] += n
                if disp=="quarantine": s["timeline"][date]["quarantine"] += n

            base = {
                "ip":ip,"count":n,"domain":domain,"org":org,
                "dkim":rec.get("dkim_result",""),"spf":rec.get("spf_result",""),
                "dmarc_result":dres,
                "header_from":rec.get("header_from",""),
                "auth_dkim_domain":rec.get("auth_dkim_domain",""),
                "auth_spf_domain":rec.get("auth_spf_domain",""),
                "auth_dkim_selector":rec.get("auth_dkim_selector",""),
                "reasons":rec.get("reasons",[]),
                "alignment":aln,"disposition":disp,
                "filename":rpt.get("filename",""),
                "ignored": ip in ignored_ips,
            }
            if disp=="quarantine": s["quarantine_records"].append(base)
            if disp=="reject":     s["reject_records"].append(base)
            if dres=="fail":       s["fail_records"].append(base)
            for se in rec.get("auth_spf_entries",[]):
                if "error" in se.get("result",""):
                    s["spf_error_records"].append({**base,
                        "spf_error":se["result"],"spf_domain":se["domain"]})

    # Serialize defaultdicts
    s["sources"]            = dict(sorted(s["sources"].items(), key=lambda x:-x[1])[:40])
    s["ip_pass"]            = dict(s["ip_pass"])
    s["ip_fail"]            = dict(s["ip_fail"])
    s["ip_quarantine"]      = dict(s["ip_quarantine"])
    s["ip_reject"]          = dict(s["ip_reject"])
    s["orgs"]               = dict(s["orgs"])
    s["domains"]            = {k: dict(v) for k,v in s["domains"].items()}
    s["dispositions"]       = dict(s["dispositions"])
    s["alignment_breakdown"]= dict(s["alignment_breakdown"])
    s["timeline"]           = {k:v for k,v in sorted(s["timeline"].items())}

    t = s["total_messages"]
    s["dmarc_pass_rate"] = round(s["dmarc_pass"]/t*100,1) if t>0 else 0
    return s


# ══════════════════════════════════════════════════════════════════
#  HTML GENERATION
# ══════════════════════════════════════════════════════════════════

def generate_html(reports, stats, ignored_ips, ip_labels):
    rj  = json.dumps(reports, ensure_ascii=False, default=str)
    sj  = json.dumps(stats,   ensure_ascii=False, default=str)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ignored_list = json.dumps(list(ignored_ips))
    ip_labels_j  = json.dumps(ip_labels)

    # Build domain list for filter
    domain_list  = sorted(stats.get("domains",{}).keys())
    domain_opts  = "\n".join(f'<option value="{d}">{d}</option>' for d in domain_list)

    return f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>DMARC Dashboard v{VERSION}</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@300;400;500;600;700&family=DM+Sans:wght@300;400;500;600;700&display=swap');
:root {{
  --bg:#07090f;--s1:#0c1120;--s2:#111827;--s3:#162032;--bd:#1f2d47;--bd2:#2d3f5a;
  --cy:#22d3ee;--gr:#4ade80;--yw:#fbbf24;--rd:#f87171;--or:#fb923c;--pu:#a78bfa;--pi:#f472b6;
  --tx:#e2e8f0;--mt:#64748b;--dm:#94a3b8;
  --mono:'IBM Plex Mono',monospace;--sans:'DM Sans',sans-serif;
  --r:14px;
}}
*,*::before,*::after{{box-sizing:border-box;margin:0;padding:0}}
html{{scroll-behavior:smooth}}
body{{font-family:var(--sans);background:var(--bg);color:var(--tx);min-height:100vh;overflow-x:hidden}}
body::after{{content:'';position:fixed;inset:0;pointer-events:none;z-index:0;
  background:radial-gradient(ellipse 80% 40% at 10% 0%,rgba(34,211,238,.04) 0%,transparent 60%),
  radial-gradient(ellipse 60% 50% at 90% 100%,rgba(74,222,128,.03) 0%,transparent 60%);}}
.w{{max-width:1750px;margin:0 auto;padding:0 28px;position:relative;z-index:1}}

/* ── HEADER ── */
header{{padding:14px 0;border-bottom:1px solid var(--bd);backdrop-filter:blur(20px);
  background:rgba(7,9,15,.92);position:sticky;top:0;z-index:400}}
.hd{{display:flex;align-items:center;justify-content:space-between;gap:16px;flex-wrap:wrap}}
.brand{{display:flex;align-items:center;gap:12px}}
.bi{{width:42px;height:42px;border-radius:10px;background:linear-gradient(135deg,var(--cy),#0891b2);
  display:flex;align-items:center;justify-content:center;font-size:20px;box-shadow:0 0 18px rgba(34,211,238,.22)}}
.bn{{font-family:var(--mono);font-size:16px;font-weight:700;color:var(--cy);letter-spacing:.08em}}
.bs{{font-family:var(--mono);font-size:9px;color:var(--mt);letter-spacing:.15em;text-transform:uppercase;margin-top:2px}}
.hr{{font-family:var(--mono);font-size:11px;color:var(--mt);text-align:right;line-height:1.8;display:flex;flex-direction:column;align-items:flex-end;gap:2px}}
.dot{{display:inline-block;width:6px;height:6px;border-radius:50%;background:var(--gr);
  box-shadow:0 0 6px var(--gr);animation:bl 2s infinite;margin-right:4px}}
@keyframes bl{{0%,100%{{opacity:1}}50%{{opacity:.3}}}}
.ver-tag{{background:rgba(34,211,238,.1);color:var(--cy);font-family:var(--mono);font-size:9px;
  padding:2px 7px;border-radius:20px;border:1px solid rgba(34,211,238,.2)}}

/* ── ALERT STRIP ── */
.asp{{margin:16px 0 0;border-radius:10px;padding:11px 16px;display:flex;align-items:center;gap:10px;
  font-size:13px;border:1px solid rgba(248,113,113,.3);
  background:linear-gradient(135deg,rgba(248,113,113,.07),rgba(248,113,113,.03));
  border-left:3px solid var(--rd);animation:fadeIn .4s ease-out}}
.asp.ok{{border-color:rgba(74,222,128,.3);border-left-color:var(--gr);
  background:linear-gradient(135deg,rgba(74,222,128,.06),rgba(74,222,128,.02))}}
@keyframes fadeIn{{from{{opacity:0;transform:translateY(-4px)}}to{{opacity:1;transform:none}}}}
.ast{{font-weight:600;color:var(--rd)}}.asp.ok .ast{{color:var(--gr)}}

/* ── MAIN ── */
main{{padding:24px 0 60px}}
section{{margin-bottom:24px}}
.sec-title{{font-family:var(--mono);font-size:10px;text-transform:uppercase;letter-spacing:.18em;
  color:var(--mt);margin-bottom:12px;display:flex;align-items:center;gap:8px}}
.sec-title::before{{content:'';display:inline-block;width:12px;height:2px;
  background:linear-gradient(90deg,var(--cy),transparent);border-radius:1px}}

/* ── KPI GRID ── */
.kr{{display:grid;grid-template-columns:repeat(auto-fill,minmax(158px,1fr));gap:12px}}
.kpi{{background:var(--s1);border:1px solid var(--bd);border-radius:var(--r);
  padding:16px 16px 12px;position:relative;overflow:hidden;
  transition:transform .15s,box-shadow .15s;cursor:pointer;user-select:none}}
.kpi:hover{{transform:translateY(-2px);box-shadow:0 8px 24px rgba(0,0,0,.45)}}
.kpi:active{{transform:translateY(0)}}
.kpi::after{{content:'';position:absolute;top:0;left:0;right:0;height:2px;background:var(--kc,var(--cy))}}
.kico{{position:absolute;right:12px;top:12px;font-size:20px;opacity:.11}}
.klb{{font-family:var(--mono);font-size:9px;text-transform:uppercase;letter-spacing:.1em;color:var(--mt);margin-bottom:5px}}
.kv{{font-family:var(--mono);font-size:30px;font-weight:700;line-height:1;color:var(--kc,var(--cy))}}
.ks{{font-size:11px;color:var(--mt);margin-top:4px;line-height:1.3}}
.kpi.good{{--kc:var(--gr)}}.kpi.warn{{--kc:var(--yw)}}.kpi.bad{{--kc:var(--rd)}}
.kpi.pu{{--kc:var(--pu)}}.kpi.pi{{--kc:var(--pi)}}

/* ── CARD ── */
.card{{background:var(--s1);border:1px solid var(--bd);border-radius:var(--r);padding:20px}}
.ch{{display:flex;align-items:center;justify-content:space-between;margin-bottom:16px;flex-wrap:wrap;gap:8px}}
.ct{{font-family:var(--mono);font-size:11px;text-transform:uppercase;letter-spacing:.1em;color:var(--dm)}}
.bge{{font-family:var(--mono);font-size:9.5px;font-weight:600;padding:3px 8px;border-radius:20px}}
.bc{{background:rgba(34,211,238,.12);color:var(--cy)}}
.bgr{{background:rgba(74,222,128,.12);color:var(--gr)}}
.by{{background:rgba(251,191,36,.12);color:var(--yw)}}
.brd{{background:rgba(248,113,113,.14);color:var(--rd)}}

/* ── GRIDS ── */
.g2{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}
.g3{{display:grid;grid-template-columns:1fr 1fr 1fr;gap:16px}}
.g31{{display:grid;grid-template-columns:2fr 1fr;gap:16px}}
.g13{{display:grid;grid-template-columns:1fr 2fr;gap:16px}}
@media(max-width:1100px){{.g2,.g3,.g31,.g13{{grid-template-columns:1fr}}}}

/* ── TABLE ── */
.tw{{overflow-x:auto;border-radius:10px}}
table{{width:100%;border-collapse:collapse;font-size:12px}}
thead tr{{border-bottom:1px solid var(--bd)}}
th{{padding:8px 12px;text-align:left;font-family:var(--mono);font-size:9px;text-transform:uppercase;
  letter-spacing:.1em;color:var(--mt);white-space:nowrap;font-weight:500;cursor:pointer;user-select:none}}
th:hover{{color:var(--cy)}}
th.sort-asc::after{{content:' ↑'}}th.sort-desc::after{{content:' ↓'}}
td{{padding:9px 12px;border-bottom:1px solid rgba(31,45,71,.5);color:var(--dm);vertical-align:middle}}
tr:last-child td{{border-bottom:none}}
tr:hover td{{background:rgba(34,211,238,.03);color:var(--tx)}}
tr.ign td{{opacity:.4;text-decoration:line-through}}
.mn{{font-family:var(--mono);font-size:11px}}
.msm{{font-family:var(--mono);font-size:10px}}

/* ── PILL ── */
.pi{{display:inline-flex;align-items:center;gap:2px;padding:2px 8px;border-radius:20px;
  font-family:var(--mono);font-size:9.5px;font-weight:600;white-space:nowrap}}
.pp{{background:rgba(74,222,128,.12);color:var(--gr);border:1px solid rgba(74,222,128,.2)}}
.pf{{background:rgba(248,113,113,.12);color:var(--rd);border:1px solid rgba(248,113,113,.2)}}
.pn{{background:rgba(100,116,139,.1);color:var(--mt);border:1px solid rgba(100,116,139,.2)}}
.pq{{background:rgba(251,191,36,.1);color:var(--yw);border:1px solid rgba(251,191,36,.2)}}
.pj{{background:rgba(248,113,113,.18);color:var(--rd);border:1px solid rgba(248,113,113,.3)}}
.ppe{{background:rgba(251,146,60,.12);color:var(--or);border:1px solid rgba(251,146,60,.2)}}
.ppu{{background:rgba(167,139,250,.12);color:var(--pu);border:1px solid rgba(167,139,250,.2)}}

/* ── TABS ── */
.tabs{{display:flex;gap:2px;border-bottom:1px solid var(--bd);margin-bottom:18px;overflow-x:auto;scrollbar-width:none}}
.tab{{padding:8px 15px;background:none;border:none;color:var(--mt);font-family:var(--mono);font-size:10px;
  cursor:pointer;border-bottom:2px solid transparent;margin-bottom:-1px;white-space:nowrap;
  letter-spacing:.06em;text-transform:uppercase;transition:color .15s,border-color .15s}}
.tab:hover{{color:var(--tx)}}.tab.active{{color:var(--cy);border-bottom-color:var(--cy)}}
.tp{{display:none}}.tp.active{{display:block}}

/* ── FILTERS ── */
.fl{{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:14px;align-items:center}}
.fi{{background:var(--s2);border:1px solid var(--bd);border-radius:8px;padding:6px 11px;
  color:var(--tx);font-family:var(--mono);font-size:10.5px;outline:none;transition:border-color .15s,box-shadow .15s}}
.fi:focus{{border-color:var(--cy);box-shadow:0 0 0 2px rgba(34,211,238,.1)}}
.fi::placeholder{{color:var(--mt)}}
select.fi option{{background:var(--s2)}}
.fi.af{{border-color:var(--cy);background:rgba(34,211,238,.06);color:var(--cy)}}
.flb{{font-family:var(--mono);font-size:10px;color:var(--mt)}}
.btn{{background:rgba(34,211,238,.08);border:1px solid var(--bd);color:var(--dm);
  font-family:var(--mono);font-size:10px;padding:6px 12px;border-radius:8px;cursor:pointer;
  transition:background .15s,color .15s}}
.btn:hover{{background:rgba(34,211,238,.15);color:var(--tx)}}
.btn.danger{{background:rgba(248,113,113,.08);border-color:rgba(248,113,113,.2);color:var(--rd)}}
.btn.danger:hover{{background:rgba(248,113,113,.18)}}

/* ── IP IGNORE PANEL ── */
.ignore-panel{{background:var(--s2);border:1px solid var(--bd);border-radius:10px;padding:16px;margin-bottom:16px}}
.ignore-panel h4{{font-family:var(--mono);font-size:11px;text-transform:uppercase;letter-spacing:.1em;
  color:var(--dm);margin-bottom:12px;display:flex;align-items:center;gap:6px}}
.ignore-input-row{{display:flex;gap:8px;flex-wrap:wrap;align-items:flex-end}}
.ignore-list{{margin-top:12px;display:flex;flex-wrap:wrap;gap:6px}}
.ignore-chip{{display:flex;align-items:center;gap:5px;background:rgba(248,113,113,.08);
  border:1px solid rgba(248,113,113,.2);border-radius:20px;padding:3px 10px;
  font-family:var(--mono);font-size:10px;color:var(--rd)}}
.chip-remove{{cursor:pointer;opacity:.6;transition:opacity .15s}}.chip-remove:hover{{opacity:1}}

/* ── PROGRESS ── */
.pg{{height:4px;background:var(--s3);border-radius:2px;overflow:hidden;margin-top:6px}}
.pgf{{height:100%;border-radius:2px}}

/* ── DOMAIN CARD ── */
.dom-card{{background:var(--s2);border:1px solid var(--bd);border-radius:10px;padding:14px;}}
.dom-card:hover{{border-color:var(--cy)}}
.dom-name{{font-family:var(--mono);font-size:14px;font-weight:700;color:var(--cy);margin-bottom:6px}}
.dom-meta{{display:flex;flex-wrap:wrap;gap:10px;font-family:var(--mono);font-size:10px;color:var(--mt)}}
.dom-grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:12px}}

/* ── REPORT CARD ── */
.rc{{background:var(--s2);border:1px solid var(--bd);border-radius:10px;padding:14px;margin-bottom:10px}}
.rc:hover{{border-color:var(--cy)}}
.rch{{display:flex;justify-content:space-between;align-items:flex-start;margin-bottom:8px;gap:10px}}
.rcn{{font-weight:600;font-size:13px}}
.rcf{{font-family:var(--mono);font-size:9.5px;color:var(--mt);margin-top:2px}}
.rm{{display:flex;flex-wrap:wrap;gap:10px;font-family:var(--mono);font-size:10px;color:var(--mt)}}

/* ── SCROLL ── */
.scr{{max-height:450px;overflow-y:auto;scrollbar-width:thin;scrollbar-color:var(--bd) transparent;padding-right:3px}}

/* ── EMPTY ── */
.emp{{text-align:center;padding:48px 20px;color:var(--mt)}}

/* ── SCROLL ANCHOR ── */
#table-section{{scroll-margin-top:76px}}

/* ── MINI STAT ── */
.mstat{{display:flex;flex-direction:column;align-items:center;padding:10px 14px;
  background:var(--s2);border-radius:8px;border:1px solid var(--bd)}}
.mstat-val{{font-family:var(--mono);font-size:22px;font-weight:700;color:var(--kc,var(--cy))}}
.mstat-lbl{{font-family:var(--mono);font-size:9px;text-transform:uppercase;letter-spacing:.08em;color:var(--mt);margin-top:3px}}

footer{{border-top:1px solid var(--bd);padding:16px 0;font-family:var(--mono);font-size:10px;color:var(--mt);
  display:flex;justify-content:space-between;flex-wrap:wrap;gap:8px}}
</style>
</head>
<body>

<header>
<div class="w">
<div class="hd">
  <div class="brand">
    <div class="bi">🛡</div>
    <div>
      <div class="bn">DMARC DASHBOARD</div>
      <div class="bs">Email Authentication Intelligence</div>
    </div>
  </div>
  <div class="hr">
    <div><span class="dot"></span><span style="color:var(--gr)">ACTIVO</span> &nbsp;<span class="ver-tag">v{VERSION}</span></div>
    <div>Generado: {now}</div>
    <div id="domain-summary" style="color:var(--cy)"></div>
  </div>
</div>
</div>
</header>

<main><div class="w">

<!-- Alert strip -->
<div id="asp" class="asp" style="display:none">
  <span id="asp-ico" style="font-size:18px">⚠️</span>
  <div><div class="ast" id="ast"></div><div style="color:var(--dm);font-size:12px;margin-top:2px" id="asb"></div></div>
</div>

<!-- KPIs -->
<section style="margin-top:18px">
  <div class="sec-title">Resumen Global <span style="font-weight:400;margin-left:6px">— clic para filtrar</span></div>
  <div class="kr" id="kpis"></div>
</section>

<!-- Charts row 1 -->
<section>
  <div class="sec-title">Autenticación</div>
  <div class="g31">
    <div class="card">
      <div class="ch"><span class="ct">Mensajes por Fecha</span><span class="bge bc" id="tlb"></span></div>
      <div style="height:220px"><canvas id="cTL"></canvas></div>
    </div>
    <div class="card">
      <div class="ch"><span class="ct">DMARC Compliance</span><span class="bge" id="gb"></span></div>
      <div style="position:relative;height:220px;display:flex;align-items:center;justify-content:center">
        <canvas id="cG" width="180" height="180" style="position:absolute"></canvas>
        <div style="text-align:center">
          <div id="gv" style="font-family:var(--mono);font-size:38px;font-weight:700"></div>
          <div style="font-family:var(--mono);font-size:9px;color:var(--mt);letter-spacing:.12em;margin-top:2px">DMARC PASS RATE</div>
          <div id="gsub" style="font-family:var(--mono);font-size:9px;color:var(--mt);margin-top:3px"></div>
        </div>
      </div>
    </div>
  </div>
</section>

<!-- Charts row 2 -->
<section>
  <div class="g2">
    <div class="card">
      <div class="ch"><span class="ct">Top IPs Origen</span><span class="bge bc" id="ipb"></span></div>
      <div style="height:250px"><canvas id="cIP"></canvas></div>
    </div>
    <div class="card">
      <div class="ch"><span class="ct">Detalle de Alineación</span></div>
      <div style="height:250px"><canvas id="cAL"></canvas></div>
    </div>
  </div>
</section>

<!-- Domains section -->
<section>
  <div class="sec-title">Dominios Monitorizados</div>
  <div class="dom-grid" id="dom-cards"></div>
</section>

<!-- Main data table -->
<section id="table-section">
  <div class="sec-title">Datos Detallados</div>
  <div class="card">
    <div class="tabs">
      <button class="tab active" data-tab="rec">Todos los Registros</button>
      <button class="tab" data-tab="alt">Alertas</button>
      <button class="tab" data-tab="ips">IPs &amp; Ignoradas</button>
      <button class="tab" data-tab="rpt">Reportes</button>
      <button class="tab" data-tab="anl">Análisis</button>
    </div>

    <!-- RECORDS TAB -->
    <div class="tp active" id="tp-rec">
      <div class="fl" id="filters-row">
        <input  class="fi" id="fip"  placeholder="🔍 IP…" style="width:150px">
        <input  class="fi" id="fhf"  placeholder="🔍 Header from…" style="width:160px">
        <select class="fi" id="fdom">
          <option value="">Dominio: Todos</option>
          {domain_opts}
        </select>
        <select class="fi" id="fdm">
          <option value="">DMARC: Todos</option>
          <option value="pass">✓ Pass</option>
          <option value="fail">✗ Fail</option>
        </select>
        <select class="fi" id="fdk">
          <option value="">DKIM: Todos</option>
          <option value="pass">✓ Pass</option>
          <option value="fail">✗ Fail</option>
        </select>
        <select class="fi" id="fsp">
          <option value="">SPF: Todos</option>
          <option value="pass">✓ Pass</option>
          <option value="fail">✗ Fail</option>
        </select>
        <select class="fi" id="fdi">
          <option value="">Disposición: Todas</option>
          <option value="none">None</option>
          <option value="quarantine">⚑ Quarantine</option>
          <option value="reject">✗ Reject</option>
        </select>
        <select class="fi" id="fal">
          <option value="">Alineación: Todas</option>
          <option value="full_pass">DKIM+SPF</option>
          <option value="dkim_only">Solo DKIM</option>
          <option value="spf_only">Solo SPF</option>
          <option value="fail">Sin alineación</option>
        </select>
        <label style="display:flex;align-items:center;gap:5px;font-family:var(--mono);font-size:10px;color:var(--mt);cursor:pointer">
          <input type="checkbox" id="fhide" style="accent-color:var(--cy)"> Ocultar ignoradas
        </label>
        <button class="btn" onclick="clearFilters()">✕ Limpiar</button>
        <span class="flb" id="rc"></span>
      </div>
      <div class="tw">
        <table id="tbl-main">
          <thead><tr>
            <th onclick="sortTable(0)">IP Origen</th>
            <th onclick="sortTable(1)">Msgs</th>
            <th onclick="sortTable(2)">Dominio</th>
            <th onclick="sortTable(3)">Header From</th>
            <th onclick="sortTable(4)">DMARC</th>
            <th onclick="sortTable(5)">DKIM</th>
            <th onclick="sortTable(6)">SPF</th>
            <th onclick="sortTable(7)">Disposición</th>
            <th onclick="sortTable(8)">Alineación</th>
            <th>Auth DKIM Domain</th>
            <th>Selector</th>
            <th>Auth SPF Domain</th>
            <th>Razones</th>
            <th>Org</th>
            <th>Archivo</th>
            <th>Acción</th>
          </tr></thead>
          <tbody id="tb"></tbody>
        </table>
      </div>
    </div>

    <!-- ALERTS TAB -->
    <div class="tp" id="tp-alt"><div id="ab"></div></div>

    <!-- IPs & IGNORADAS TAB -->
    <div class="tp" id="tp-ips">
      <!-- Ignore panel -->
      <div class="ignore-panel">
        <h4>⊘ Gestión de IPs Ignoradas</h4>
        <div class="ignore-input-row">
          <div>
            <div style="font-family:var(--mono);font-size:9px;color:var(--mt);margin-bottom:4px">IP o rango</div>
            <input class="fi" id="ig-ip" placeholder="192.168.1.1" style="width:180px">
          </div>
          <div>
            <div style="font-family:var(--mono);font-size:9px;color:var(--mt);margin-bottom:4px">Etiqueta</div>
            <input class="fi" id="ig-label" placeholder="Descripción..." style="width:200px">
          </div>
          <button class="btn danger" onclick="addIgnoredIP()">+ Añadir a ignoradas</button>
          <span style="font-family:var(--mono);font-size:9px;color:var(--mt);align-self:flex-end;padding-bottom:7px">
            También puedes editar directamente la hoja "IPs Ignoradas" del Excel
          </span>
        </div>
        <div class="ignore-list" id="ignore-list"></div>
      </div>

      <!-- IPs table -->
      <div class="tw">
        <table>
          <thead><tr>
            <th>IP</th><th>Total msgs</th><th>DMARC Pass</th><th>DMARC Fail</th>
            <th>% Pass</th><th>Cuarentena</th><th>Rechazados</th>
            <th>Dominios</th><th>Etiqueta</th><th>Ignorada</th>
          </tr></thead>
          <tbody id="ip-tbody"></tbody>
        </table>
      </div>
    </div>

    <!-- REPORTS TAB -->
    <div class="tp" id="tp-rpt"><div class="scr" id="rb"></div></div>

    <!-- ANALYSIS TAB -->
    <div class="tp" id="tp-anl">
      <div class="g3" style="margin-bottom:18px" id="anl-mini-stats"></div>
      <div class="g2">
        <div class="card" style="background:var(--s2)">
          <div class="ch"><span class="ct">Tendencia Semanal</span></div>
          <div style="height:200px"><canvas id="cWeek"></canvas></div>
        </div>
        <div class="card" style="background:var(--s2)">
          <div class="ch"><span class="ct">Orgs Reportadoras</span><span class="bge bc" id="orgs-count"></span></div>
          <div style="height:200px"><canvas id="cOrgs"></canvas></div>
        </div>
      </div>
      <div style="margin-top:16px" class="card" style="background:var(--s2)">
        <div class="ch"><span class="ct">Top Senders con Fallos</span></div>
        <div class="tw" id="fail-senders"></div>
      </div>
    </div>
  </div>
</section>

</div></main>

<footer><div class="w" style="display:flex;justify-content:space-between;width:100%;flex-wrap:wrap;gap:8px">
  <span>DMARC DASHBOARD v{VERSION} // dmarc_visualizer.py</span>
  <span id="fr"></span>
</div></footer>

<script>
const REPORTS     = JSON.parse({repr(rj)});
const STATS       = JSON.parse({repr(sj)});
const INIT_IGNORED= new Set(JSON.parse({repr(ignored_list)}));
const INIT_LABELS = JSON.parse({repr(ip_labels_j)});

let ignoredIPs = new Set(INIT_IGNORED);
let ipLabels   = {{...INIT_LABELS}};

const $ = id => document.getElementById(id);
const n = v  => (+(v||0)).toLocaleString('es-ES');
const pct = (a,b) => b>0 ? (a/b*100).toFixed(1)+'%' : '—';

// ── PILL helpers ──────────────────────────────────────────────────
function pill(v) {{
  if(!v||v==='—') return '<span class="pi pn">—</span>';
  const c=v.toLowerCase();
  const m={{pass:'pp',fail:'pf',quarantine:'pq',reject:'pj',permerror:'ppe',temperror:'ppe'}};
  const cls=m[c]||'pn';
  const ic=c==='pass'?'✓ ':c==='fail'?'✗ ':c==='quarantine'?'⚑ ':c==='reject'?'⛔ ':'';
  return `<span class="pi ${{cls}}">${{ic}}${{v}}</span>`;
}}
function apill(a) {{
  const m={{full_pass:['pp','✓ DKIM+SPF'],dkim_only:['ppu','◑ Solo DKIM'],
            spf_only:['ppu','◑ Solo SPF'],fail:['pf','✗ Ninguno']}};
  const [cls,lb]=m[a]||['pn',a||'—'];
  return `<span class="pi ${{cls}}">${{lb}}</span>`;
}}

// ── ALERT STRIP ───────────────────────────────────────────────────
(()=>{{
  const qr=(STATS.quarantine_records||[]);
  const rr=(STATS.reject_records||[]);
  const fr=(STATS.fail_records||[]).filter(r=>r.disposition==='none');
  const parts=[];
  const qsum=qr.reduce((a,x)=>a+x.count,0);
  const rsum=rr.reduce((a,x)=>a+x.count,0);
  if(rsum>0) parts.push(`${{n(rsum)}} rechazados`);
  if(qsum>0) parts.push(`${{n(qsum)}} en cuarentena`);
  if(STATS.dmarc_fail>0) parts.push(`${{n(STATS.dmarc_fail)}} DMARC fail`);
  const el=$('asp'); el.style.display='flex';
  if(parts.length){{
    $('ast').textContent='⚠ Problemas detectados en los reportes';
    $('asb').textContent=parts.join(' · ');
  }} else {{
    el.classList.add('ok'); $('asp-ico').textContent='✅';
    $('ast').textContent='✓ Sin alertas — autenticación DMARC correcta';
  }}
}})();

// ── DOMAIN SUMMARY in header ──────────────────────────────────────
(()=>{{
  const doms=Object.keys(STATS.domains||{{}});
  $('domain-summary').textContent=doms.length+' dominio'+(doms.length!==1?'s':'')+': '+doms.join(', ');
}})();

// ── KPIs ──────────────────────────────────────────────────────────
(()=>{{
  const pr=STATS.dmarc_pass_rate;
  const ab=STATS.alignment_breakdown||{{}};
  const qsum=(STATS.quarantine_records||[]).reduce((a,x)=>a+x.count,0);
  const rsum=(STATS.reject_records||[]).reduce((a,x)=>a+x.count,0);
  const kpis=[
    {{l:'Total Mensajes',  v:n(STATS.total_messages),  s:STATS.total_reports+' reporte(s)',                cls:'',    ico:'📧', f:null}},
    {{l:'DMARC Pass Rate', v:pr+'%',                    s:'(DKIM ó SPF alineado)',                         cls:pr>=90?'good':pr>=70?'warn':'bad', ico:'📊', f:null}},
    {{l:'DMARC Pass',      v:n(STATS.dmarc_pass),       s:'mensajes autenticados',                         cls:'good',ico:'✅', f:{{e:'fdm',v:'pass'}}}},
    {{l:'DMARC Fail',      v:n(STATS.dmarc_fail),       s:'sin autenticación válida',                     cls:STATS.dmarc_fail>0?'bad':'', ico:'🔴', f:{{e:'fdm',v:'fail'}}}},
    {{l:'Solo DKIM',       v:n(ab.dkim_only||0),        s:'SPF no alineado (OK DMARC)',                   cls:'pu',  ico:'🔑', f:{{e:'fal',v:'dkim_only'}}}},
    {{l:'Solo SPF',        v:n(ab.spf_only||0),         s:'DKIM no alineado (OK DMARC)',                  cls:'pu',  ico:'📮', f:{{e:'fal',v:'spf_only'}}}},
    {{l:'Cuarentena',      v:n(qsum),                   s:'mensajes en quarantine',                        cls:qsum>0?'bad':'',  ico:'⚑',  f:{{e:'fdi',v:'quarantine'}}}},
    {{l:'Rechazados',      v:n(rsum),                   s:'mensajes en reject',                            cls:rsum>0?'bad':'',  ico:'🚫', f:{{e:'fdi',v:'reject'}}}},
    {{l:'IPs únicas',      v:n(Object.keys(STATS.sources||{{}}).length), s:'fuentes detectadas',           cls:'',    ico:'🌐', f:null}},
    {{l:'Dominios',        v:n(Object.keys(STATS.domains||{{}}).length), s:'dominios monitorizados',       cls:'pi',  ico:'🏷', f:null}},
  ];
  $('kpis').innerHTML=kpis.map((k,i)=>
    `<div class="kpi ${{k.cls}}" onclick="kpiClick(${{JSON.stringify(k.f)}})">
      <div class="kico">${{k.ico}}</div>
      <div class="klb">${{k.l}}</div>
      <div class="kv">${{k.v}}</div>
      <div class="ks">${{k.s}}</div>
    </div>`
  ).join('');
}})();

function kpiClick(f) {{
  if(!f) return;
  $(f.e).value=f.v; $(f.e).classList.add('af');
  document.querySelectorAll('.tab').forEach(t=>t.classList.remove('active'));
  document.querySelectorAll('.tp').forEach(t=>t.classList.remove('active'));
  document.querySelector('[data-tab="rec"]').classList.add('active');
  $('tp-rec').classList.add('active');
  filterTable();
  setTimeout(()=>$('table-section').scrollIntoView({{behavior:'smooth',block:'start'}}),80);
}}

// ── CHARTS ────────────────────────────────────────────────────────
Chart.defaults.color='#64748b';Chart.defaults.font.family="'IBM Plex Mono',monospace";Chart.defaults.font.size=10;
const gd={{color:'rgba(31,45,71,.6)'}};

// Timeline
(()=>{{
  const tl=STATS.timeline||{{}};const lbs=Object.keys(tl);
  if(!lbs.length){{$('tlb').textContent='SIN DATOS';return;}}
  $('tlb').textContent=lbs.length>1?lbs[0]+' → '+lbs[lbs.length-1]:lbs[0];
  new Chart($('cTL').getContext('2d'),{{type:'bar',data:{{labels:lbs,datasets:[
    {{label:'DMARC Pass',data:lbs.map(d=>tl[d].pass||0),backgroundColor:'rgba(74,222,128,.55)',borderColor:'rgba(74,222,128,.9)',borderWidth:1,borderRadius:4}},
    {{label:'DMARC Fail',data:lbs.map(d=>tl[d].fail||0),backgroundColor:'rgba(248,113,113,.5)',borderColor:'rgba(248,113,113,.8)',borderWidth:1,borderRadius:4}},
    {{label:'Cuarentena',data:lbs.map(d=>tl[d].quarantine||0),backgroundColor:'rgba(251,191,36,.4)',borderColor:'rgba(251,191,36,.7)',borderWidth:1,borderRadius:4}},
  ]}},options:{{responsive:true,maintainAspectRatio:false,
    plugins:{{legend:{{labels:{{color:'#94a3b8',boxWidth:10}}}}}},
    scales:{{x:{{grid:gd}},y:{{grid:gd,beginAtZero:true,ticks:{{precision:0}}}}}}}}}});
}})();

// Gauge
(()=>{{
  const r=STATS.dmarc_pass_rate;
  const col=r>=90?'#4ade80':r>=70?'#fbbf24':'#f87171';
  const bc=r>=90?'bgr':r>=70?'by':'brd';
  $('gb').textContent=r>=90?'✓ EXCELENTE':r>=70?'~ ACEPTABLE':'✗ CRÍTICO';
  $('gb').className='bge '+bc;
  $('gv').textContent=r+'%';$('gv').style.color=col;
  $('gsub').textContent=n(STATS.dmarc_pass)+' pass / '+n(STATS.dmarc_fail)+' fail';
  new Chart($('cG').getContext('2d'),{{type:'doughnut',
    data:{{datasets:[{{data:[r,100-r],backgroundColor:[col,'rgba(22,32,50,.8)'],borderWidth:0,circumference:270,rotation:225}}]}},
    options:{{cutout:'76%',responsive:false,plugins:{{legend:{{display:false}},tooltip:{{enabled:false}}}}}}}});
}})();

// IP bar
(()=>{{
  const src=STATS.sources||{{}};const ips=Object.keys(src).slice(0,15);
  $('ipb').textContent=Object.keys(src).length+' IPs';
  new Chart($('cIP').getContext('2d'),{{type:'bar',
    data:{{labels:ips,datasets:[{{label:'Msgs',data:ips.map(ip=>src[ip]),
      backgroundColor:ips.map((_,i)=>`hsla(${{200-i*11}},76%,60%,.65)`),borderRadius:3,borderWidth:0}}]}},
    options:{{indexAxis:'y',responsive:true,maintainAspectRatio:false,
      plugins:{{legend:{{display:false}}}},
      scales:{{x:{{grid:gd,beginAtZero:true}},y:{{grid:{{display:false}},ticks:{{font:{{size:9}}}}}}}}}}}});
}})();

// Alignment donut
(()=>{{
  const ab=STATS.alignment_breakdown||{{}};
  const total=STATS.total_messages||1;
  new Chart($('cAL').getContext('2d'),{{type:'doughnut',
    data:{{
      labels:[
        `DKIM+SPF (${{Math.round((ab.full_pass||0)/total*100)}}%)`,
        `Solo DKIM (${{Math.round((ab.dkim_only||0)/total*100)}}%)`,
        `Solo SPF (${{Math.round((ab.spf_only||0)/total*100)}}%)`,
        `Sin alineación (${{Math.round((ab.fail||0)/total*100)}}%)`,
      ],
      datasets:[{{data:[ab.full_pass||0,ab.dkim_only||0,ab.spf_only||0,ab.fail||0],
        backgroundColor:['rgba(74,222,128,.75)','rgba(167,139,250,.7)','rgba(34,211,238,.65)','rgba(248,113,113,.65)'],
        borderColor:'rgba(7,9,15,.6)',borderWidth:2,hoverOffset:8}}]}},
    options:{{responsive:true,maintainAspectRatio:false,
      plugins:{{legend:{{position:'bottom',labels:{{color:'#94a3b8',padding:10,boxWidth:10,font:{{size:10}}}}}}}}}}}});
}})();

// ── DOMAIN CARDS ──────────────────────────────────────────────────
(()=>{{
  const doms=STATS.domains||{{}};
  const html=Object.entries(doms).map(([dom,ds])=>{{
    const msgs=ds.msgs||0,pass=ds.pass||0,fail=ds.fail||0;
    const pr=msgs>0?Math.round(pass/msgs*100):0;
    const col=pr>=90?'var(--gr)':pr>=70?'var(--yw)':'var(--rd)';
    return`<div class="dom-card">
      <div class="dom-name">${{dom}}</div>
      <div class="dom-meta">
        <span>📧 ${{n(msgs)}} msgs</span>
        <span style="color:var(--gr)">✓ ${{n(pass)}} pass</span>
        <span style="color:var(--rd)">✗ ${{n(fail)}} fail</span>
        <span style="color:var(--yw)">⚑ ${{n(ds.quar||0)}} cuar</span>
        <span style="color:var(--rd)">⛔ ${{n(ds.rej||0)}} rej</span>
      </div>
      <div class="pg"><div class="pgf" style="width:${{pr}}%;background:${{col}}"></div></div>
      <div style="font-family:var(--mono);font-size:10px;color:var(--mt);margin-top:4px">
        ${{pr}}% DMARC pass
        <button class="btn" style="margin-left:8px;padding:2px 8px;font-size:9px"
          onclick="filterByDomain('${{dom}}')">🔍 Filtrar</button>
      </div>
    </div>`;
  }}).join('');
  $('dom-cards').innerHTML=html||'<div class="emp">Sin dominios</div>';
}})();

function filterByDomain(dom) {{
  $('fdom').value=dom; $('fdom').classList.add('af');
  document.querySelectorAll('.tab').forEach(t=>t.classList.remove('active'));
  document.querySelectorAll('.tp').forEach(t=>t.classList.remove('active'));
  document.querySelector('[data-tab="rec"]').classList.add('active');
  $('tp-rec').classList.add('active');
  filterTable();
  setTimeout(()=>$('table-section').scrollIntoView({{behavior:'smooth',block:'start'}}),80);
}}

// ══════════════════════════════════════════════════════════════════
//  RECORDS TABLE
// ══════════════════════════════════════════════════════════════════
let allRows=[], sortCol=-1, sortDir=1;

(()=>{{
  REPORTS.forEach(r=>{{
    (r.records||[]).forEach(rec=>allRows.push({{
      ...rec,_org:r.org_name||'?',_file:r.filename||'',_domain:r.domain||'?'
    }}));
  }});
  renderRecs(allRows);
}})();

function renderRecs(rows) {{
  const hideIgnored=$('fhide').checked;
  $('rc').textContent=n(rows.length)+' registros';
  const tb=$('tb');
  const visible=rows.filter(r=>!(hideIgnored && ignoredIPs.has(r.source_ip)));
  if(!visible.length){{
    tb.innerHTML='<tr><td colspan="16" class="emp">Sin resultados</td></tr>'; return;
  }}
  tb.innerHTML=visible.map(r=>{{
    const ignored=ignoredIPs.has(r.source_ip);
    const lbl=ipLabels[r.source_ip]||'';
    const rs=(r.reasons||[]).map(x=>x.type+(x.comment?':'+x.comment:'')).filter(Boolean).join('; ');
    return`<tr class="${{ignored?'ign':''}}">
      <td class="mn">${{r.source_ip||'—'}}${{lbl?` <span style="color:var(--yw);font-size:9px">[${{lbl}}]</span>`:''}}</td>
      <td class="mn" style="color:var(--cy);font-weight:600">${{r.count||0}}</td>
      <td class="msm" style="color:var(--cy)">${{r._domain||'—'}}</td>
      <td class="msm">${{r.header_from||'—'}}</td>
      <td>${{pill(r.dmarc_result)}}</td>
      <td>${{pill(r.dkim_result)}}</td>
      <td>${{pill(r.spf_result)}}</td>
      <td>${{pill(r.disposition)}}</td>
      <td>${{apill(r.alignment)}}</td>
      <td class="msm" style="color:var(--mt)">${{r.auth_dkim_domain||'—'}}</td>
      <td class="msm" style="color:var(--mt);font-size:9px">${{r.auth_dkim_selector||'—'}}</td>
      <td class="msm" style="color:var(--mt)">${{r.auth_spf_domain||'—'}}</td>
      <td class="msm" style="font-size:9px;color:var(--mt);max-width:120px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis" title="${{rs}}">${{rs||'—'}}</td>
      <td class="msm" style="color:var(--mt);font-size:9px">${{r._org}}</td>
      <td class="msm" style="color:var(--mt);font-size:9px;max-width:100px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis">${{r._file}}</td>
      <td>
        <button style="background:${{ignored?'rgba(100,116,139,.1)':'rgba(248,113,113,.08)'}};
          border:1px solid ${{ignored?'rgba(100,116,139,.2)':'rgba(248,113,113,.2)'}};
          color:${{ignored?'var(--mt)':'var(--rd)'}};font-family:var(--mono);font-size:9px;
          padding:2px 7px;border-radius:6px;cursor:pointer"
          onclick="toggleIgnore('${{r.source_ip.replace(/'/g,'')}}',this)">
          ${{ignored?'↩':'⊘'}}
        </button>
      </td>
    </tr>`;
  }}).join('');
}}

// Table sort
function sortTable(col) {{
  const ths=document.querySelectorAll('#tbl-main th');
  ths.forEach((t,i)=>{{t.classList.remove('sort-asc','sort-desc');if(i===col)t.classList.add(sortDir>0?'sort-asc':'sort-desc');}});
  if(sortCol===col) sortDir*=-1; else {{sortCol=col;sortDir=1;}}
  const keys=['source_ip','count','_domain','header_from','dmarc_result','dkim_result','spf_result','disposition','alignment'];
  const k=keys[col]||'source_ip';
  allRows.sort((a,b)=>{{
    const av=a[k]||'',bv=b[k]||'';
    if(typeof av==='number'||typeof bv==='number') return sortDir*((+av||0)-(+bv||0));
    return sortDir*String(av).localeCompare(String(bv));
  }});
  filterTable();
}}

// Filters
['fip','fhf','fdom','fdm','fdk','fsp','fdi','fal','fhide'].forEach(id=>{{
  const el=$(id); if(!el) return;
  el.addEventListener(id==='fhide'?'change':'input',()=>{{
    el.value ? el.classList.add('af') : el.classList.remove('af');
    filterTable();
  }});
}});

function filterTable() {{
  const ip  =$('fip').value.toLowerCase();
  const hf  =$('fhf').value.toLowerCase();
  const dom =$('fdom').value;
  const dm  =$('fdm').value;
  const dk  =$('fdk').value;
  const sp  =$('fsp').value;
  const di  =$('fdi').value;
  const al  =$('fal').value;
  const filtered=allRows.filter(r=>
    (!ip  ||(r.source_ip||'').toLowerCase().includes(ip))&&
    (!hf  ||(r.header_from||'').toLowerCase().includes(hf))&&
    (!dom ||r._domain===dom)&&
    (!dm  ||r.dmarc_result===dm)&&
    (!dk  ||r.dkim_result===dk)&&
    (!sp  ||r.spf_result===sp)&&
    (!di  ||r.disposition===di)&&
    (!al  ||r.alignment===al)
  );
  renderRecs(filtered);
}}

function clearFilters() {{
  ['fip','fhf','fdom','fdm','fdk','fsp','fdi','fal'].forEach(id=>{{
    const el=$(id); el.value=''; el.classList.remove('af');
  }});
  $('fhide').checked=false;
  filterTable();
}}

// ── IGNORE SYSTEM ─────────────────────────────────────────────────
function renderIgnoreList() {{
  const list=$('ignore-list');
  if(ignoredIPs.size===0){{list.innerHTML='<span style="font-family:var(--mono);font-size:10px;color:var(--mt)">Sin IPs ignoradas</span>';return;}}
  list.innerHTML=[...ignoredIPs].map(ip=>
    `<div class="ignore-chip">
      ${{ipLabels[ip]?`<span style="color:var(--yw)">${{ipLabels[ip]}}</span> — `:''}}${{ip}}
      <span class="chip-remove" onclick="removeIgnoredIP('${{ip}}')">✕</span>
    </div>`
  ).join('');
}}

function addIgnoredIP() {{
  const ip=$('ig-ip').value.trim();
  const lbl=$('ig-label').value.trim();
  if(!ip) return;
  ignoredIPs.add(ip);
  if(lbl) ipLabels[ip]=lbl;
  $('ig-ip').value=''; $('ig-label').value='';
  renderIgnoreList();
  renderIPTable();
  filterTable();
}}

function removeIgnoredIP(ip) {{
  ignoredIPs.delete(ip);
  delete ipLabels[ip];
  renderIgnoreList();
  renderIPTable();
  filterTable();
}}

function toggleIgnore(ip, btn) {{
  if(ignoredIPs.has(ip)) removeIgnoredIP(ip);
  else {{ ignoredIPs.add(ip); renderIgnoreList(); renderIPTable(); filterTable(); }}
  // re-render happens in filterTable
}}

// ── IP TABLE ──────────────────────────────────────────────────────
function renderIPTable() {{
  const src=STATS.sources||{{}};
  const ipDoms={{}};
  REPORTS.forEach(r=>{{
    (r.records||[]).forEach(rec=>{{
      if(!ipDoms[rec.source_ip]) ipDoms[rec.source_ip]=new Set();
      ipDoms[rec.source_ip].add(r.domain||'?');
    }});
  }});

  const tb=$('ip-tbody');
  tb.innerHTML=Object.entries(src).map(([ip,tot])=>{{
    const pass=(STATS.ip_pass||{{}})[ip]||0;
    const fail=(STATS.ip_fail||{{}})[ip]||0;
    const quar=(STATS.ip_quarantine||{{}})[ip]||0;
    const rej =(STATS.ip_reject||{{}})[ip]||0;
    const pr  =tot>0?Math.round(pass/tot*100):0;
    const doms=[...(ipDoms[ip]||[])].join(', ');
    const lbl =ipLabels[ip]||'';
    const ign =ignoredIPs.has(ip);
    return`<tr class="${{ign?'ign':''}}">
      <td class="mn">${{ip}}</td>
      <td class="mn" style="color:var(--cy)">${{n(tot)}}</td>
      <td class="mn" style="color:var(--gr)">${{n(pass)}}</td>
      <td class="mn" style="color:${{fail>0?'var(--rd)':'var(--mt)'}}">${{n(fail)}}</td>
      <td><span class="mn" style="color:${{pr>=90?'var(--gr)':pr>=70?'var(--yw)':'var(--rd)'}}">${{pr}}%</span></td>
      <td class="mn" style="color:${{quar>0?'var(--yw)':'var(--mt)'}}">${{n(quar)}}</td>
      <td class="mn" style="color:${{rej>0?'var(--rd)':'var(--mt)'}}">${{n(rej)}}</td>
      <td class="msm" style="color:var(--mt);font-size:9.5px;max-width:180px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis" title="${{doms}}">${{doms}}</td>
      <td class="msm" style="color:var(--yw)">${{lbl||'—'}}</td>
      <td>
        <button style="background:${{ign?'rgba(100,116,139,.1)':'rgba(248,113,113,.08)'}};
          border:1px solid ${{ign?'rgba(100,116,139,.2)':'rgba(248,113,113,.2)'}};
          color:${{ign?'var(--mt)':'var(--rd)'}};font-family:var(--mono);font-size:9px;
          padding:2px 7px;border-radius:6px;cursor:pointer"
          onclick="toggleIgnore('${{ip}}',this)">
          ${{ign?'↩ Restaurar':'⊘ Ignorar'}}
        </button>
      </td>
    </tr>`;
  }}).join('');
}}

// ── ALERTS TAB ────────────────────────────────────────────────────
(()=>{{
  const sec=(title,icon,rows,col)=>{{
    if(!rows.length) return '';
    const total=rows.reduce((a,x)=>a+x.count,0);
    return`<div style="margin-bottom:22px">
      <div style="display:flex;align-items:center;gap:8px;margin-bottom:10px">
        <span style="font-size:16px">${{icon}}</span>
        <span style="font-weight:600;color:${{col}}">${{title}}</span>
        <span class="bge" style="background:rgba(0,0,0,.2);color:${{col}}">${{rows.length}} registros · ${{n(total)}} msgs</span>
      </div>
      <div class="tw"><table>
        <thead><tr><th>IP</th><th>Msgs</th><th>Dominio</th><th>Header From</th>
          <th>DMARC</th><th>DKIM</th><th>SPF</th><th>Disposición</th>
          <th>Auth DKIM</th><th>Auth SPF</th><th>Razones</th><th>Ignorada</th></tr></thead>
        <tbody>${{rows.map(r=>`<tr class="${{r.ignored?'ign':''}}">
          <td class="mn">${{r.ip}}</td>
          <td class="mn" style="color:${{col}};font-weight:700">${{r.count}}</td>
          <td class="msm" style="color:var(--cy)">${{r.domain}}</td>
          <td class="msm">${{r.header_from||'—'}}</td>
          <td>${{pill(r.dmarc_result)}}</td><td>${{pill(r.dkim)}}</td><td>${{pill(r.spf)}}</td>
          <td>${{pill(r.disposition)}}</td>
          <td class="msm" style="color:var(--mt)">${{r.auth_dkim_domain||'—'}}</td>
          <td class="msm" style="color:var(--mt)">${{r.auth_spf_domain||'—'}}</td>
          <td class="msm" style="font-size:9px;color:var(--mt)">${{(r.reasons||[]).map(x=>x.type).join(', ')||'—'}}</td>
          <td style="font-family:var(--mono);font-size:9px;color:${{r.ignored?'var(--yw)':'var(--mt)'}}">${{r.ignored?'Sí':'—'}}</td>
        </tr>`).join('')}}</tbody>
      </table></div>
    </div>`;
  }};
  let html='';
  html+=sec('Rechazados (reject)','🚫',STATS.reject_records||[],'#f87171');
  html+=sec('Cuarentena','⚑',STATS.quarantine_records||[],'#fbbf24');
  html+=sec('DMARC Fail sin disposición especial','🔴',(STATS.fail_records||[]).filter(r=>r.disposition==='none'),'#f87171');
  html+=sec('Errores SPF','⚙',STATS.spf_error_records||[],'#fb923c');
  $('ab').innerHTML=html||'<div class="emp" style="color:var(--gr);font-size:16px;font-weight:600">✅ Sin alertas</div>';
  const tot=(STATS.quarantine_records||[]).length+(STATS.reject_records||[]).length+(STATS.fail_records||[]).filter(r=>r.disposition==='none').length;
  if(tot) document.querySelector('[data-tab="alt"]').textContent=`⚠ Alertas (${{tot}})`;
}})();

// ── REPORTS TAB ───────────────────────────────────────────────────
(()=>{{
  const pol=p=>p==='r'?'relaxed':p==='s'?'strict':p||'?';
  $('rb').innerHTML=REPORTS.map(r=>{{
    const s=r.summary||{{}};const pr=s.dmarc_pass_rate||0;
    const bc=pr>=90?'bgr':pr>=70?'by':'brd';
    const col=pr>=90?'var(--gr)':pr>=70?'var(--yw)':'var(--rd)';
    const isNew=r._is_new;
    return`<div class="rc" style="${{isNew?'border-color:rgba(74,222,128,.3);':''}}">
      <div class="rch">
        <div>
          <div class="rcn">
            ${{isNew?'<span style="background:rgba(74,222,128,.15);color:var(--gr);font-family:var(--mono);font-size:9px;padding:1px 7px;border-radius:10px;margin-right:6px">NEW</span>':''}}
            ${{r.org_name||'?'}} — <span style="color:var(--cy)">${{r.domain||'?'}}</span>
          </div>
          <div class="rcf">📄 ${{r.filename}} · ID: ${{r.report_id||'?'}}</div>
        </div>
        <span class="bge ${{bc}}">${{pr}}% DMARC pass</span>
      </div>
      <div class="rm">
        <span>📅 ${{r.date_begin||'?'}} → ${{r.date_end||'?'}}</span>
        <span>📧 ${{n(s.total_messages||0)}} msgs</span>
        <span>📋 ${{s.record_count||0}} recs</span>
        <span>🌐 ${{s.unique_ips||0}} IPs</span>
        <span style="color:var(--gr)">✓ ${{n(s.dmarc_pass||0)}} pass</span>
        <span style="color:var(--rd)">✗ ${{n(s.dmarc_fail||0)}} fail</span>
        <span style="color:var(--pu)">◑ DKIM ${{n(s.dkim_only||0)}} / SPF ${{n(s.spf_only||0)}}</span>
        <span style="color:var(--yw)">⚑ ${{n(s.quarantined||0)}}</span>
        <span style="color:var(--rd)">⛔ ${{n(s.rejected||0)}}</span>
        <span>🔑 DKIM:${{pol(r.policy_adkim)}} 📮 SPF:${{pol(r.policy_aspf)}} 🛡 P:${{r.policy_p||'?'}}/SP:${{r.policy_sp||'?'}}</span>
      </div>
      <div class="pg"><div class="pgf" style="width:${{pr}}%;background:linear-gradient(90deg,${{col}},rgba(0,0,0,.3))"></div></div>
    </div>`;
  }}).join('')||'<div class="emp">Sin reportes cargados</div>';
}})();

// ── ANALYSIS TAB ──────────────────────────────────────────────────
(()=>{{
  // Mini stats
  const ab=STATS.alignment_breakdown||{{}};
  const total=STATS.total_messages||1;
  const miniStats=[
    {{l:'Full DKIM+SPF', v:n(ab.full_pass||0), sub:Math.round((ab.full_pass||0)/total*100)+'%', kc:'var(--gr)'}},
    {{l:'Solo DKIM (ok)', v:n(ab.dkim_only||0), sub:Math.round((ab.dkim_only||0)/total*100)+'%', kc:'var(--pu)'}},
    {{l:'Solo SPF (ok)',  v:n(ab.spf_only||0),  sub:Math.round((ab.spf_only||0)/total*100)+'%', kc:'var(--cy)'}},
    {{l:'Sin alineación',v:n(ab.fail||0),        sub:Math.round((ab.fail||0)/total*100)+'%',     kc:'var(--rd)'}},
    {{l:'En cuarentena', v:n((STATS.quarantine_records||[]).reduce((a,x)=>a+x.count,0)), sub:'msgs', kc:'var(--yw)'}},
    {{l:'Rechazados',    v:n((STATS.reject_records||[]).reduce((a,x)=>a+x.count,0)),     sub:'msgs', kc:'var(--rd)'}},
  ];
  $('anl-mini-stats').innerHTML=miniStats.map(m=>
    `<div class="mstat" style="--kc:${{m.kc}}">
      <div class="mstat-val">${{m.v}}</div>
      <div class="mstat-lbl">${{m.l}} <span style="color:var(--cy)">${{m.sub}}</span></div>
    </div>`
  ).join('');

  // Weekly trend (group by week)
  const tl=STATS.timeline||{{}};
  const weeks={{}};
  Object.entries(tl).forEach(([date,d])=>{{
    const dt=new Date(date);
    const mon=new Date(dt);mon.setDate(dt.getDate()-dt.getDay()+1);
    const wk=mon.toISOString().slice(0,10);
    if(!weeks[wk]) weeks[wk]={{pass:0,fail:0}};
    weeks[wk].pass+=d.pass||0;weeks[wk].fail+=d.fail||0;
  }});
  const wlbs=Object.keys(weeks).sort();
  new Chart($('cWeek').getContext('2d'),{{type:'line',
    data:{{labels:wlbs,datasets:[
      {{label:'Pass',data:wlbs.map(w=>weeks[w].pass),borderColor:'#4ade80',backgroundColor:'rgba(74,222,128,.1)',tension:.3,fill:true,pointRadius:3}},
      {{label:'Fail',data:wlbs.map(w=>weeks[w].fail),borderColor:'#f87171',backgroundColor:'rgba(248,113,113,.1)',tension:.3,fill:true,pointRadius:3}},
    ]}},
    options:{{responsive:true,maintainAspectRatio:false,
      plugins:{{legend:{{labels:{{color:'#94a3b8',boxWidth:10}}}}}},
      scales:{{x:{{grid:gd}},y:{{grid:gd,beginAtZero:true,ticks:{{precision:0}}}}}}}}}});

  // Orgs chart
  const orgs=STATS.orgs||{{}};
  const orgNames=Object.keys(orgs).slice(0,10);
  $('orgs-count').textContent=Object.keys(orgs).length+' orgs';
  new Chart($('cOrgs').getContext('2d'),{{type:'bar',
    data:{{labels:orgNames,datasets:[{{label:'Reportes',data:orgNames.map(o=>orgs[o]),
      backgroundColor:'rgba(34,211,238,.55)',borderRadius:3,borderWidth:0}}]}},
    options:{{indexAxis:'y',responsive:true,maintainAspectRatio:false,
      plugins:{{legend:{{display:false}}}},
      scales:{{x:{{grid:gd,beginAtZero:true,ticks:{{precision:0}}}},y:{{grid:{{display:false}},ticks:{{font:{{size:9}}}}}}}}}}}});

  // Top fail senders
  const failRows=(STATS.fail_records||[])
    .filter(r=>r.disposition==='none')
    .sort((a,b)=>b.count-a.count).slice(0,15);
  $('fail-senders').innerHTML=failRows.length?
    `<table><thead><tr><th>IP</th><th>Msgs</th><th>Dominio</th><th>Header From</th><th>DKIM</th><th>SPF</th><th>Auth DKIM Domain</th></tr></thead>
    <tbody>${{failRows.map(r=>`<tr>
      <td class="mn">${{r.ip}}</td>
      <td class="mn" style="color:var(--rd);font-weight:600">${{r.count}}</td>
      <td class="msm" style="color:var(--cy)">${{r.domain}}</td>
      <td class="msm">${{r.header_from||'—'}}</td>
      <td>${{pill(r.dkim)}}</td><td>${{pill(r.spf)}}</td>
      <td class="msm" style="color:var(--mt)">${{r.auth_dkim_domain||'—'}}</td>
    </tr>`).join('')}}</tbody></table>`:
    '<div class="emp" style="color:var(--gr)">✓ Sin fallos sin disposición</div>';
}})();

// ── TABS ──────────────────────────────────────────────────────────
document.querySelectorAll('.tab').forEach(t=>{{
  t.addEventListener('click',()=>{{
    document.querySelectorAll('.tab').forEach(x=>x.classList.remove('active'));
    document.querySelectorAll('.tp').forEach(x=>x.classList.remove('active'));
    t.classList.add('active');
    $('tp-'+t.dataset.tab).classList.add('active');
  }});
}});

// ── INIT ──────────────────────────────────────────────────────────
renderIgnoreList();
renderIPTable();
$('fr').textContent=`${{STATS.total_reports}} reportes · ${{n(STATS.total_messages)}} mensajes · ${{Object.keys(STATS.domains||{{}}).join(', ')}}`;
</script>
</body>
</html>"""


# ══════════════════════════════════════════════════════════════════
#  SERVER + BROWSER
# ══════════════════════════════════════════════════════════════════

def find_port(start=8765):
    for p in range(start, start+100):
        try:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                s.bind(("127.0.0.1", p)); return p
        except OSError: continue
    return None


def serve_and_open(html_path):
    html_path = Path(html_path).resolve()
    port = find_port()
    if not port:
        print(f"  ⚠  Sin puerto libre. Abre: file://{html_path}"); return

    class H(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *a, **kw):
            super().__init__(*a, directory=str(html_path.parent), **kw)
        def log_message(self, *a): pass

    srv = http.server.HTTPServer(("127.0.0.1", port), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{port}/{html_path.name}"
    print(f"\n  🌐 Servidor : {url}")
    print(f"  🔥 Abriendo navegador...\n")
    time.sleep(0.3)
    webbrowser.open(url)
    print("  ✅ Dashboard activo. Presiona Ctrl+C para salir.\n")
    try:
        while True: time.sleep(1)
    except KeyboardInterrupt:
        print("\n  Cerrando..."); srv.shutdown()


# ══════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════

def main():
    args    = sys.argv[1:]
    folder  = "./dmarc_reports"
    output  = "./dmarc_dashboard.html"
    no_open = False
    for a in args:
        if a == "--no-open": no_open = True
        elif a.endswith(".html"): output = a
        else: folder = a

    print()
    print("━" * 62)
    print(f"   DMARC REPORT VISUALIZER  v{VERSION}")
    print("━" * 62)

    # ── Excel ────────────────────────────────────────────────────
    wb, ignored_ips, ip_labels, seen_hashes = init_excel(folder)

    # ── Load reports ─────────────────────────────────────────────
    all_reports, new_reports = load_folder(folder, seen_hashes)

    if not all_reports:
        print("  ⚠  No hay reportes que mostrar.")
        print(f"     Coloca archivos .xml/.zip/.gz en: {Path(folder).resolve()}")
        sys.exit(0)

    # Mark new reports for display
    new_hashes = {r["_hash"] for r in new_reports}
    for r in all_reports:
        r["_is_new"] = r["_hash"] in new_hashes

    # ── Stats ────────────────────────────────────────────────────
    stats = compute_stats(all_reports, ignored_ips, ip_labels)

    print(f"\n  Estadísticas globales:")
    print(f"    Mensajes totales  : {stats['total_messages']}")
    print(f"    DMARC pass rate   : {stats['dmarc_pass_rate']}%")
    print(f"    DMARC pass        : {stats['dmarc_pass']}")
    print(f"    DMARC fail        : {stats['dmarc_fail']}")
    print(f"    Cuarentena msgs   : {sum(r['count'] for r in stats.get('quarantine_records',[]))}")
    print(f"    Rechazados msgs   : {sum(r['count'] for r in stats.get('reject_records',[]))}")
    print(f"    Dominios          : {', '.join(stats['domains'].keys())}")
    print(f"    Nuevos reportes   : {len(new_reports)}")
    print(f"    IPs ignoradas     : {len(ignored_ips)}")

    # ── Update Excel ─────────────────────────────────────────────
    update_excel(wb, all_reports, stats, folder, new_reports)

    # ── Generate HTML ─────────────────────────────────────────────
    html = generate_html(all_reports, stats, ignored_ips, ip_labels)
    out  = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(html, encoding="utf-8")
    print(f"\n  💾 Dashboard : {out.resolve()}")
    print(f"  📊 Excel     : {(Path(folder)/EXCEL_FILE).resolve()}")
    print("━" * 62)

    if no_open: print(f"\n  Abre manualmente: file://{out.resolve()}")
    else: serve_and_open(out)


if __name__ == "__main__":
    main()